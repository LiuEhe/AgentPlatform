import asyncio
import os
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlsplit
from uuid import uuid4

import pytest
from docx import Document

from app.schemas.conversation_turn import ConversationTurnInput, TurnMessage


@pytest.fixture(autouse=True)
def isolate_external_settings(monkeypatch, tmp_path):
    import os

    for key in tuple(os.environ):
        if key.startswith(("CONVERSATION_", "LANGFUSE_")):
            monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("LANGFUSE_ENABLED", "false")
    monkeypatch.setenv("LANGFUSE_PROMPTS_ENABLED", "false")
    monkeypatch.setenv("DOCUMENTS_DIR", str(tmp_path / "documents"))
    for key in (
        "LLM_PROVIDER",
        "LLM_MODEL",
        "LLM_API_KEY",
        "LLM_BASE_URL",
        "RAGFLOW_BASE_URL",
        "RAGFLOW_API_KEY",
        "RAGFLOW_DATASET_ID",
        "INSTITUTION_REVIEW_CONCURRENCY",
        "DOCUMENT_EXPORT_TIMEOUT_SECONDS",
        "DOCUMENT_EXPORT_CONCURRENCY",
        "DOCUMENT_EXPORT_MAX_CHARS",
        "DOCUMENT_EXPORT_MAX_BODY_BYTES",
        "DOCUMENT_EXPORT_MAX_OUTPUT_BYTES",
        "DATABASE_URL",
    ):
        monkeypatch.delenv(key, raising=False)


@pytest.fixture
def postgres_target():
    import asyncpg

    dsn = os.environ.get("TEST_DATABASE_URL")
    if not dsn or urlsplit(dsn).path != "/agent_platform_test":
        pytest.fail("真实数据库测试需要显式 TEST_DATABASE_URL，库名必须为 agent_platform_test")
    schema = "history_test_" + uuid4().hex

    async def prepare():
        conn = await asyncpg.connect(dsn)
        try:
            await conn.execute(f'CREATE SCHEMA "{schema}"')
            await conn.execute(f'SET search_path TO "{schema}"')
            await conn.execute(
                (
                    Path(__file__).parents[1] / "database/schema_v2.sql"
                ).read_text(encoding="utf-8")
            )
        finally:
            await conn.close()

    async def cleanup():
        conn = await asyncpg.connect(dsn)
        try:
            # schema 来自本次 uuid4，不接收外部标识符，不清理应用/public schema。
            await conn.execute(f'DROP SCHEMA "{schema}" CASCADE')
        finally:
            await conn.close()

    asyncio.run(prepare())
    yield dsn, schema
    asyncio.run(cleanup())


@pytest.fixture
def fake_model_factory():
    class FakeModel:
        def __init__(self, handler=None):
            self.calls = []
            self.handler = handler

        async def call(self, schema, messages):
            self.calls.append((schema, messages))
            if self.handler:
                return await self.handler(schema, messages)
            return schema() if schema else "未发现术语问题，无需基于本次检查修改。"

        def with_structured_output(self, schema):
            async def invoke(messages):
                return await self.call(schema, messages)

            return SimpleNamespace(ainvoke=invoke)

        async def ainvoke(self, messages):
            return await self.call(None, messages)

    return FakeModel


@pytest.fixture
def docx_bytes_factory():
    def make(*, title="草稿", body="第一条 正文"):
        doc = Document()
        doc.add_heading(title, level=1)
        for line in body.splitlines():
            doc.add_paragraph(line)
        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    return make


@pytest.fixture
def turn_factory():
    def make(*, role="user", content="你好", documents=(), continuation=None, history=()):
        return ConversationTurnInput(
            latest_message_id="m1",
            messages=(
                *history,
                TurnMessage(
                    id="m1",
                    role=role,
                    content=content,
                    document_ids=tuple(d.id for d in documents),
                ),
            ),
            documents=documents,
            continuation=continuation,
        )

    return make


@pytest.fixture(autouse=True)
def wire_application_database(monkeypatch, request):
    if "create_app" not in request.module.__dict__:
        return
    postgres_target = request.getfixturevalue("postgres_target")
    from pydantic import SecretStr

    import app.main as main_module
    from app.infrastructure.conversation.postgres_store import PostgresConversationStore

    dsn, schema = postgres_target

    def make_store(settings, **kwargs):
        settings = settings.model_copy(update={"database_url": SecretStr(dsn)})
        return PostgresConversationStore(settings, _schema=schema, **kwargs)

    monkeypatch.setattr(main_module, "PostgresConversationStore", make_store)
