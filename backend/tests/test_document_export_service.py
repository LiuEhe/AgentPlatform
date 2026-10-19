import asyncio
import json
import threading
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.core.config import Settings
from app.core.errors import AppError
from app.services import document_export_service as exports


class FakeRunner:
    def __init__(self, content, ast=None):
        self.content = content
        self.ast = ast or {'pandoc-api-version': [1, 23, 1], 'meta': {}, 'blocks': []}
        self.calls = []
        self.closed = False

    async def initialize(self):
        pass

    async def run(self, data, **options):
        self.calls.append((data, options))
        if options['target_format'] == 'json':
            return json.dumps(self.ast).encode()
        return self.content

    async def aclose(self):
        self.closed = True


@pytest.fixture
def rich_docx():
    doc = Document()
    doc.add_heading('制度审核报告', 0)
    paragraph = doc.add_paragraph('中文正文 ')
    paragraph.add_run('重点').bold = True
    paragraph.add_run('说明').italic = True
    doc.add_paragraph('第一项', style='List Number')
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), 'heading')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '链接'
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '条款'
    table.cell(0, 1).text = '依据'
    table.cell(1, 0).text = '第一条'
    table.cell(1, 1).text = '完整依据' * 1000 + '依据尾部'
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def perform_export(runner, markdown='## 正文', title='制度审核报告', **config):
    async def scenario():
        service = exports.DocumentExportService(Settings(_env_file=None, **config), runner)
        await service.initialize()
        return await service.export_docx(markdown, title=title)
    return asyncio.run(scenario())


def test_exports_complete_chinese_report(rich_docx):
    runner = FakeRunner(rich_docx)
    result = perform_export(runner)
    doc = Document(BytesIO(result.content))
    assert result.filename == '制度审核报告.docx'
    assert result.media_type.endswith('wordprocessingml.document')
    assert doc.tables[0].cell(1, 1).text.endswith('依据尾部')
    assert runner.calls[0][1]['deadline'] == runner.calls[1][1]['deadline']


@pytest.mark.parametrize('markdown,title,config', [
    ('  ', '标题', {}), ('正文', '\r\n', {}), ('太长了', '标题',
                                               {'document_export_max_chars': 2}),
])
def test_direct_service_input_limits_match_http(rich_docx, markdown, title, config):
    runner = FakeRunner(rich_docx)
    with pytest.raises(AppError) as raised:
        perform_export(runner, markdown, title, **config)
    assert raised.value.code == 'INVALID_EXPORT_INPUT' and not runner.calls


@pytest.mark.parametrize('kind', ['Image', 'RawInline', 'RawBlock'])
def test_image_and_raw_nodes_are_rejected(rich_docx, kind):
    runner = FakeRunner(rich_docx, {
        'pandoc-api-version': [1], 'meta': {}, 'blocks': [{'t': kind, 'c': []}],
    })
    with pytest.raises(AppError) as raised:
        perform_export(runner)
    assert raised.value.code == 'EXPORT_CONTENT_UNSUPPORTED'
    assert len(runner.calls) == 1


def test_code_with_image_or_html_text_is_preserved(rich_docx):
    runner = FakeRunner(rich_docx, {
        'pandoc-api-version': [1], 'meta': {},
        'blocks': [{'t': 'CodeBlock', 'c': [['', [], []], '![图](x) <b>标签</b>']}],
    })
    assert perform_export(runner).content


def test_word_layout_preserves_runs_numbering_and_links(rich_docx):
    doc = Document(BytesIO(perform_export(FakeRunner(rich_docx)).content))
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21
    assert round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 1) == 2.5
    assert doc.styles['Normal'].font.size.pt == 11
    assert doc.styles['Normal'].paragraph_format.line_spacing == 1.5
    assert doc.styles['Title'].font.size.pt == 18
    assert any(run.bold for p in doc.paragraphs for run in p.runs)
    assert any(run.italic for p in doc.paragraphs for run in p.runs)
    assert any(p.style.name == 'List Number' for p in doc.paragraphs)
    assert doc.element.body.xpath('.//w:hyperlink')
    assert doc.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert doc.tables[0]._tbl.xpath('./w:tblPr/w:tblBorders')


@pytest.mark.parametrize('title,filename', [
    ('a/b\\c', 'abc.docx'), ('CON', 'review-report.docx'),
    ('<>:"/\\|?*', 'review-report.docx'), ('报告.docx', '报告.docx'),
    ('x' * 200, 'x' * 120 + '.docx'),
])
def test_filename_cannot_be_path_or_reserved_name(rich_docx, title, filename):
    assert perform_export(FakeRunner(rich_docx), title=title).filename == filename


def test_same_leading_title_is_not_duplicated(rich_docx):
    runner = FakeRunner(rich_docx, {
        'pandoc-api-version': [1], 'meta': {}, 'blocks': [
            {'t': 'Header', 'c': [1, ['', [], []], [{'t': 'Str', 'c': '制度审核报告'}]]},
        ],
    })
    perform_export(runner)
    ast = json.loads(runner.calls[1][0])
    assert 'title' not in ast['meta']


def test_invalid_docx_returns_conversion_failure():
    with pytest.raises(AppError) as raised:
        perform_export(FakeRunner(b'not a docx'))
    assert raised.value.code == 'DOCUMENT_EXPORT_FAILED'


class BlockingRunner(FakeRunner):
    def __init__(self, content, target=1):
        super().__init__(content)
        self.release = asyncio.Event()
        self.started = asyncio.Event()
        self.active = self.maximum = self.entered = 0
        self.target = target

    async def run(self, data, **options):
        if options['source_format'] == 'gfm':
            self.active += 1
            self.entered += 1
            self.maximum = max(self.maximum, self.active)
            if self.active == self.target:
                self.started.set()
            try:
                await self.release.wait()
            finally:
                self.active -= 1
        return await super().run(data, **options)


def test_requests_share_one_concurrency_limit(rich_docx):
    async def scenario():
        runner = BlockingRunner(rich_docx, target=2)
        service = exports.DocumentExportService(Settings(_env_file=None), runner)
        await service.initialize()
        tasks = [asyncio.create_task(service.export_docx('正文')) for _ in range(3)]
        await asyncio.wait_for(runner.started.wait(), 2)
        await asyncio.sleep(0)
        assert runner.maximum == runner.entered == 2
        runner.release.set()
        assert len(await asyncio.wait_for(asyncio.gather(*tasks), 3)) == 3
        assert runner.maximum == 2
    asyncio.run(scenario())


def test_queue_timeout_returns_busy_without_spawning(monkeypatch, rich_docx):
    async def scenario():
        monkeypatch.setattr(exports, 'EXPORT_QUEUE_TIMEOUT_SECONDS', .01, raising=False)
        runner = BlockingRunner(rich_docx)
        service = exports.DocumentExportService(
            Settings(_env_file=None, document_export_concurrency=1), runner,
        )
        await service.initialize()
        first = asyncio.create_task(service.export_docx('正文'))
        await asyncio.wait_for(runner.started.wait(), 2)
        try:
            with pytest.raises(AppError) as raised:
                await service.export_docx('正文')
            assert raised.value.code == 'DOCUMENT_EXPORT_BUSY'
            assert runner.entered == 1
        finally:
            runner.release.set()
            await first
    asyncio.run(scenario())


def test_cancel_queued_request_does_not_release_another_slot(rich_docx):
    async def scenario():
        runner = BlockingRunner(rich_docx)
        service = exports.DocumentExportService(
            Settings(_env_file=None, document_export_concurrency=1), runner,
        )
        await service.initialize()
        first = asyncio.create_task(service.export_docx('正文'))
        await asyncio.wait_for(runner.started.wait(), 2)
        queued = asyncio.create_task(service.export_docx('正文'))
        await asyncio.sleep(0)
        queued.cancel()
        with pytest.raises(asyncio.CancelledError):
            await queued
        third = asyncio.create_task(service.export_docx('正文'))
        await asyncio.sleep(0)
        assert runner.entered == 1
        runner.release.set()
        await asyncio.wait_for(asyncio.gather(first, third), 3)
    asyncio.run(scenario())


def test_cancel_during_formatting_keeps_slot_until_thread_finishes(monkeypatch, rich_docx):
    async def scenario():
        entered, release = threading.Event(), threading.Event()
        original = exports._format_docx
        def slow(content):
            entered.set()
            assert release.wait(3)
            return original(content)
        monkeypatch.setattr(exports, '_format_docx', slow)
        runner = FakeRunner(rich_docx)
        service = exports.DocumentExportService(
            Settings(_env_file=None, document_export_concurrency=1), runner,
        )
        await service.initialize()
        first = asyncio.create_task(service.export_docx('正文'))
        try:
            assert await asyncio.to_thread(entered.wait, 2)
            first.cancel()
            await asyncio.sleep(0)
            second = asyncio.create_task(service.export_docx('正文'))
            await asyncio.sleep(0)
            assert not first.done() and len(runner.calls) == 2
        finally:
            release.set()
        with pytest.raises(asyncio.CancelledError):
            await first
        assert (await asyncio.wait_for(second, 3)).content
    asyncio.run(scenario())


def test_failed_export_does_not_cancel_other_request(rich_docx):
    async def scenario():
        runner = FakeRunner(rich_docx)
        original = runner.run
        async def fail(data, **options):
            if data == b'fail':
                raise AppError('DOCUMENT_EXPORT_FAILED', '失败', 500)
            return await original(data, **options)
        runner.run = fail
        service = exports.DocumentExportService(Settings(_env_file=None), runner)
        await service.initialize()
        results = await asyncio.gather(
            service.export_docx('fail'), service.export_docx('成功'), return_exceptions=True,
        )
        assert isinstance(results[0], AppError) and results[1].content
    asyncio.run(scenario())


def test_close_reaps_active_exports_and_is_idempotent(rich_docx):
    async def scenario():
        runner = BlockingRunner(rich_docx)
        service = exports.DocumentExportService(Settings(_env_file=None), runner)
        await service.initialize()
        task = asyncio.create_task(service.export_docx('正文'))
        await asyncio.wait_for(runner.started.wait(), 2)
        await service.aclose()
        await service.aclose()
        assert task.cancelled() and runner.closed
        with pytest.raises(AppError) as raised:
            await service.export_docx('正文')
        assert raised.value.status_code == 503
    asyncio.run(scenario())


@pytest.mark.parametrize('ast', [[], {'blocks': 'bad', 'meta': {}, 'pandoc-api-version': [1]}])
def test_malformed_ast_returns_conversion_failure(rich_docx, ast):
    runner = FakeRunner(rich_docx)
    runner.ast = ast
    with pytest.raises(AppError) as raised:
        perform_export(runner)
    assert raised.value.code == 'DOCUMENT_EXPORT_FAILED'


def test_final_output_is_not_truncated_when_over_limit(rich_docx):
    with pytest.raises(AppError) as raised:
        perform_export(FakeRunner(rich_docx), document_export_max_output_bytes=500)
    assert raised.value.code == 'EXPORT_OUTPUT_TOO_LARGE'
