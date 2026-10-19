import asyncio
from io import BytesIO
from urllib.parse import unquote

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.core.errors import AppError
from app.main import create_app
from app.schemas.document_export import ExportedDocument


@pytest.fixture
def fake_export(monkeypatch, docx_bytes_factory):
    class Export:
        def __init__(self):
            self.calls = []
            self.initialized = self.closed = 0
            self.failure = None

        async def initialize(self):
            self.initialized += 1

        async def export_docx(self, markdown, *, title):
            self.calls.append((markdown, title))
            if self.failure:
                raise self.failure
            return ExportedDocument(docx_bytes_factory(), '制度报告.docx')

        async def aclose(self):
            self.closed += 1

    service = Export()
    monkeypatch.setattr('app.main.DocumentExportService', lambda *_: service, raising=False)
    return service


def config(**changes):
    return Settings(_env_file=None, llm_model=None, llm_api_key=None, **changes)


def test_docx_download_works_without_model_configuration(fake_export):
    with TestClient(create_app(config())) as client:
        result = client.post('/api/v1/documents/export/docx', json={'markdown': '报告'})
        assert result.status_code == 200
        assert Document(BytesIO(result.content)).paragraphs
        assert result.headers['content-type'].endswith('wordprocessingml.document')
        assert result.headers['cache-control'] == 'no-store'
        assert client.app.state.agent_runtime is None
    assert fake_export.calls == [('报告', '制度审核报告')]


def test_custom_api_prefix_is_used(fake_export):
    with TestClient(create_app(config(api_prefix='/custom'))) as client:
        response = client.post('/custom/documents/export/docx', json={'markdown': '文'})
        assert response.status_code == 200


@pytest.mark.parametrize('content,content_type', [
    ('{', 'application/json'), ('[]', 'application/json'),
    ('{"markdown": ""}', 'application/json'),
    ('{"markdown": "文", "path": "private"}', 'application/json'),
    ('{"markdown": "文", "title": "a\\r\\nb"}', 'application/json'),
    ('{"markdown": "文"}', 'text/plain'),
])
def test_malformed_or_extra_fields_return_422_before_conversion(
    fake_export, content, content_type,
):
    with TestClient(create_app(config())) as client:
        response = client.post('/api/v1/documents/export/docx', content=content.encode(),
                               headers={'Content-Type': content_type})
    assert response.status_code == 422
    assert response.json()['error']['code'] == 'INVALID_EXPORT_INPUT'
    assert not fake_export.calls


@pytest.mark.parametrize('declared', [None, '1'])
def test_chunked_body_limit_without_trusting_content_length(fake_export, declared):
    async def scenario():
        app = create_app(config(document_export_max_body_bytes=32))
        async def chunks():
            yield b'{"markdown":"'
            yield b'x' * 40
            yield b'"}'
        headers = {'Content-Type': 'application/json'}
        if declared:
            headers['Content-Length'] = declared
        async with app.router.lifespan_context(app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
                                         base_url='http://test') as client:
                response = await client.post('/api/v1/documents/export/docx',
                                             content=chunks(), headers=headers)
        assert response.status_code == 413
        assert response.json()['error']['code'] == 'EXPORT_REQUEST_TOO_LARGE'
        assert not fake_export.calls
    asyncio.run(scenario())


def test_filename_header_is_utf8_and_has_no_injection(fake_export):
    with TestClient(create_app(config())) as client:
        response = client.post('/api/v1/documents/export/docx', json={'markdown': '文'})
    header = response.headers['content-disposition']
    assert 'filename="review-report.docx"' in header
    assert unquote(header.split("filename*=UTF-8''")[1]) == '制度报告.docx'
    assert '\r' not in header and '\n' not in header


def test_openapi_documents_json_and_binary_response(fake_export):
    with TestClient(create_app(config())) as client:
        schema = client.get('/openapi.json').json()
    operation = schema['paths']['/api/v1/documents/export/docx']['post']
    assert 'application/json' in operation['requestBody']['content']
    content = operation['responses']['200']['content']
    assert next(iter(content.values()))['schema']['format'] == 'binary'


def test_lifespan_shares_service_and_closes_once(fake_export):
    with TestClient(create_app(config())) as client:
        for _ in range(2):
            assert client.post('/api/v1/documents/export/docx',
                               json={'markdown': '文'}).status_code == 200
        assert fake_export.initialized == 1
    assert fake_export.closed == 1


def test_error_response_and_logs_omit_private_data(fake_export, caplog):
    fake_export.failure = AppError('DOCUMENT_EXPORT_FAILED', '文档转换失败', 500)
    with TestClient(create_app(config())) as client:
        response = client.post('/api/v1/documents/export/docx',
                               json={'markdown': 'private body', 'title': 'private title'})
    assert response.status_code == 500
    assert response.json()['error']['request_id']
    assert 'private body' not in response.text + caplog.text
    assert 'private title' not in response.text + caplog.text


def test_missing_bundled_binary_only_disables_export(monkeypatch, fake_model_factory):
    async def unavailable(_):
        raise AppError('DOCUMENT_EXPORT_UNAVAILABLE', '文档转换程序不可用', 503)
    monkeypatch.setattr('app.main.PandocRunner.initialize', unavailable)
    monkeypatch.setattr('app.main.create_llm_service', lambda _: fake_model_factory())
    settings = Settings(_env_file=None, llm_model='test', llm_api_key='test',
                        ragflow_base_url=None, ragflow_api_key=None, ragflow_dataset_id=None)
    with TestClient(create_app(settings)) as client:
        assert client.get('/health/ready').status_code == 200
        assert client.app.state.agent_runtime is not None
        response = client.post('/api/v1/documents/export/docx', json={'markdown': '文'})
        assert response.status_code == 503
        assert response.json()['error']['code'] == 'DOCUMENT_EXPORT_UNAVAILABLE'
