from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document

from app.nodes.document.parse import DocumentParseError, parse_docx


def test_parse_preserves_headings_and_tables():
    document = Document()
    document.add_heading("草稿", 1)
    document.add_paragraph("第一条 原文")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "栏目"
    table.cell(1, 0).text = "条款"
    buffer = BytesIO()
    document.save(buffer)
    result = parse_docx(buffer.getvalue(), max_chars=50000)
    assert "# 草稿" in result and "第一条 原文" in result and "| 栏目" in result


@pytest.mark.parametrize("content", [b"", b"broken"])
def test_invalid_docx_has_stable_code(content):
    with pytest.raises(DocumentParseError) as raised:
        parse_docx(content, max_chars=50000)
    assert raised.value.code == "INVALID_DOCX"


def test_archive_entry_limit():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for index in range(2001):
            archive.writestr(str(index), b"")
    with pytest.raises(DocumentParseError) as raised:
        parse_docx(buffer.getvalue(), max_chars=50000)
    assert raised.value.code == "UPLOAD_TOO_LARGE"


def test_declared_uncompressed_limit(monkeypatch):
    class Archive:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def infolist(self):
            return [SimpleNamespace(file_size=104857601)]

    monkeypatch.setattr("app.nodes.document.parse.ZipFile", lambda _: Archive())
    with pytest.raises(DocumentParseError) as raised:
        parse_docx(b"not parsed", max_chars=50000)
    assert raised.value.code == "UPLOAD_TOO_LARGE"


def test_markdown_length_limit(docx_bytes_factory):
    with pytest.raises(DocumentParseError) as raised:
        parse_docx(docx_bytes_factory(body="很长的正文"), max_chars=2)
    assert raised.value.code == "REVIEW_CONTENT_TOO_LONG"
