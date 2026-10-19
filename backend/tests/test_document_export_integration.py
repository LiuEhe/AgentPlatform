"""真实包内 Pandoc 集成测试，不调用模型、检索或远程资源。"""

import asyncio
import subprocess
import sys
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO

import httpx
import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from uvicorn import Config

from app.core.config import Settings
from app.infrastructure.pandoc import PandocRunner, resolve_bundled_pandoc
from app.main import create_app


@pytest.fixture(autouse=True)
def require_bundled_pandoc():
    try:
        version('pypandoc-binary')
    except PackageNotFoundError:
        pytest.skip('真实转换需要已声明的 pypandoc-binary 依赖')


def settings():
    return Settings(_env_file=None, llm_model=None, llm_api_key=None)


def test_real_export_keeps_all_four_sections_and_long_evidence():
    evidence = '这是标准文档中需要完整保留的中文依据。' * 250 + '唯一依据尾标'
    source = (
        '## 一. 简述\n\n中文报告 **重点** 和 *强调*。\n\n'
        '## 二. 冲突检测\n\n'
        '| 条款 | 问题 | 说明 | 程度 | 来源 | 依据 |\n'
        '| --- | --- | --- | --- | --- | --- |\n'
        f'| 第一条 | 冲突 | 原文说明 | 高 | 标准 | {evidence} |\n\n'
        '## 三. 术语检测\n\n1. 第一项\n   - 嵌套项目\n2. 第二项\n\n'
        '## 四. 修改建议\n\n[标准链接](https://example.com/standard)\n\n'
        '```text\n![示例](local.png) <b>原样代码</b>\n```\n\n'
        + '\n\n'.join('#' * level + f' 层级{level}' for level in range(1, 7))
    )
    with TestClient(create_app(settings())) as client:
        response = client.post('/api/v1/documents/export/docx',
                               json={'markdown': source, 'title': '测试报告'})
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    doc = Document(BytesIO(response.content))
    text = '\n'.join(p.text for p in doc.paragraphs)
    for part in ('测试报告', '一. 简述', '二. 冲突检测', '三. 术语检测', '四. 修改建议',
                 '嵌套项目', '原样代码'):
        assert part in text
    assert doc.tables[0].cell(1, 5).text == evidence
    assert doc.element.body.xpath('.//w:numPr')
    assert doc.element.body.xpath('.//w:hyperlink')
    assert any(run.bold for p in doc.paragraphs for run in p.runs)
    assert any(run.italic for p in doc.paragraphs for run in p.runs)
    assert doc.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    fonts = doc.styles['Normal']._element.rPr.rFonts
    assert fonts.get(qn('w:eastAsia')) == '宋体'
    assert round(doc.sections[0].page_width.cm, 1) == 21
    assert sum(c.width for c in doc.tables[0].columns) <= (
        doc.sections[0].page_width - doc.sections[0].left_margin
        - doc.sections[0].right_margin + 10000
    )


@pytest.mark.parametrize('source', ['![图](private-file.png)', '<b>原始HTML</b>'])
def test_real_engine_rejects_media_but_keeps_code_literals(source):
    with TestClient(create_app(settings())) as client:
        response = client.post('/api/v1/documents/export/docx', json={'markdown': source})
        assert response.status_code == 422
        assert response.json()['error']['code'] == 'EXPORT_CONTENT_UNSUPPORTED'
        literal = client.post('/api/v1/documents/export/docx',
                              json={'markdown': f'```\n{source}\n```'})
        assert literal.status_code == 200
        assert source in '\n'.join(p.text for p in Document(BytesIO(literal.content)).paragraphs)


@pytest.mark.parametrize('heading', ['制度审核报告', '**制度审核报告**', '*制度审核报告*'])
def test_same_leading_title_is_not_duplicated(heading):
    with TestClient(create_app(settings())) as client:
        response = client.post('/api/v1/documents/export/docx',
                               json={'markdown': f'# {heading}\n\n正文'})
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    assert [p.text for p in doc.paragraphs].count('制度审核报告') == 1
    assert doc.paragraphs[0].style.name == 'Title'


def test_table_inline_code_keeps_monospace_font():
    with TestClient(create_app(settings())) as client:
        response = client.post('/api/v1/documents/export/docx',
                               json={'markdown': '| code |\n| --- |\n| `a + b` |'})
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    run = next(run for p in doc.tables[0].cell(1, 0).paragraphs for run in p.runs
               if run.text == 'a + b')
    assert run.font.name == 'Consolas'


def test_installed_distribution_is_used_for_both_steps(monkeypatch):
    async def scenario():
        binary = resolve_bundled_pandoc()
        runner = PandocRunner(settings())
        await runner.initialize()
        spawn = subprocess.Popen
        commands = []
        def record(args, **options):
            commands.append(tuple(args))
            return spawn(args, **options)
        monkeypatch.setattr(subprocess, 'Popen', record)
        deadline = asyncio.get_running_loop().time() + 5
        ast = await runner.run(b'# text', source_format='gfm', target_format='json',
                               deadline=deadline)
        result = await runner.run(ast, source_format='json', target_format='docx',
                                  deadline=deadline)
        await runner.aclose()
        assert Document(BytesIO(result)).paragraphs
        assert len(commands) == 2
        assert all(c[0] == str(binary) and '--sandbox' in c for c in commands)
    asyncio.run(scenario())


@pytest.mark.skipif(sys.platform != 'win32', reason='Windows reload 事件循环回归')
def test_real_http_export_on_windows_reload_event_loop():
    async def scenario():
        app = create_app(settings())
        async with app.router.lifespan_context(app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
                                         base_url='http://test') as client:
                response = await client.post('/api/v1/documents/export/docx',
                                             json={'markdown': '## 中文审核报告\n\n正文'})
        assert response.status_code == 200
        assert Document(BytesIO(response.content)).paragraphs
    server_config = Config('app.main:create_app', factory=True, reload=True)
    with asyncio.Runner(loop_factory=server_config.get_loop_factory()) as runner:
        runner.run(scenario())
