import asyncio
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from app.core.config import Settings
from app.core.errors import AppError
from app.services.review_service import ReviewService


class FakeRuntime:
    def __init__(self, output: object = "# 审核报告") -> None:
        self.output = output
        self.calls: list[tuple[str, dict[str, object]]] = []

    async def invoke(self, agent_id: str, payload: dict[str, object]) -> SimpleNamespace:
        self.calls.append((agent_id, payload))
        return SimpleNamespace(output=self.output)


def make_docx(*, title: str = "经营管理制度", body: str = "第一条 适用范围") -> bytes:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(body)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_service(
    *,
    max_upload_bytes: int = 20_971_520,
    max_review_chars: int = 50_000,
    runtime: FakeRuntime | None = None,
) -> tuple[ReviewService, FakeRuntime]:
    fake_runtime = runtime or FakeRuntime()
    settings = Settings(
        _env_file=None,
        max_upload_bytes=max_upload_bytes,
        max_review_chars=max_review_chars,
    )
    return ReviewService(settings, fake_runtime), fake_runtime


def run_review(service: ReviewService, filename: str, content: bytes) -> str:
    return asyncio.run(service.review_document(filename, content))


def test_rejects_non_docx_upload() -> None:
    service, runtime = make_service()

    with pytest.raises(AppError) as raised:
        run_review(service, "policy.pdf", b"content")

    assert raised.value.code == "UNSUPPORTED_DOCUMENT_TYPE"
    assert raised.value.status_code == 415
    assert runtime.calls == []


def test_rejects_empty_document() -> None:
    service, runtime = make_service()

    with pytest.raises(AppError) as raised:
        run_review(service, "policy.docx", b"")

    assert raised.value.code == "EMPTY_DOCUMENT"
    assert raised.value.status_code == 422
    assert runtime.calls == []


def test_rejects_upload_over_configured_byte_limit() -> None:
    service, runtime = make_service(max_upload_bytes=8)

    with pytest.raises(AppError) as raised:
        run_review(service, "policy.docx", b"123456789")

    assert raised.value.code == "UPLOAD_TOO_LARGE"
    assert raised.value.status_code == 413
    assert runtime.calls == []


def test_rejects_unreadable_docx() -> None:
    service, runtime = make_service()

    with pytest.raises(AppError) as raised:
        run_review(service, "policy.docx", b"not a Word document")

    assert raised.value.code == "INVALID_DOCX"
    assert raised.value.status_code == 422
    assert runtime.calls == []


def test_rejects_converted_document_over_character_limit() -> None:
    service, runtime = make_service(max_review_chars=1_000)

    with pytest.raises(AppError) as raised:
        run_review(service, "policy.docx", make_docx(body="条款" * 600))

    assert raised.value.code == "REVIEW_CONTENT_TOO_LONG"
    assert raised.value.status_code == 422
    assert runtime.calls == []


def test_converts_document_and_returns_agent_report() -> None:
    service, runtime = make_service()

    report = run_review(service, "fallback-name.docx", make_docx())

    assert report == "# 审核报告"
    assert runtime.calls == [
        (
            "institution_review",
            {
                "title": "经营管理制度",
                "content": "# 经营管理制度\n\n第一条 适用范围",
            },
        )
    ]


def test_standard_mode_converts_both_documents_and_preserves_roles() -> None:
    service, runtime = make_service()
    report = asyncio.run(service.review_document(
        "同名.docx", make_docx(title="草稿", body="第一条 草稿正文"),
        review_mode="standard_document", standard_filename="同名.docx",
        standard_content=make_docx(title="标准", body="必须复核"),
    ))
    assert report == "# 审核报告"
    payload = runtime.calls[0][1]
    assert payload["content"] == "# 草稿\n\n第一条 草稿正文"
    assert payload["standard_content"] == "# 标准\n\n必须复核"
    assert payload["standard_name"] == "同名.docx"
    assert payload["review_mode"] == "standard_document"


@pytest.mark.parametrize("filename,content,expected", [
    ("standard.pdf", b"x", "UNSUPPORTED_DOCUMENT_TYPE"),
    ("standard.docx", b"", "EMPTY_DOCUMENT"),
    ("standard.docx", b"not docx", "INVALID_DOCX"),
])
def test_rejects_invalid_standard_before_runtime(filename, content, expected) -> None:
    service, runtime = make_service()
    with pytest.raises(AppError) as raised:
        asyncio.run(service.review_document(
            "草稿.docx", make_docx(), review_mode="standard_document",
            standard_filename=filename, standard_content=content,
        ))
    assert raised.value.code == expected
    assert "标准" in raised.value.message
    assert runtime.calls == []


def test_rejects_standard_character_limit_before_runtime() -> None:
    service, runtime = make_service(max_review_chars=1000)
    with pytest.raises(AppError) as raised:
        asyncio.run(service.review_document(
            "草稿.docx", make_docx(), review_mode="standard_document",
            standard_filename="标准.docx", standard_content=make_docx(body="标准" * 600),
        ))
    assert raised.value.code == "REVIEW_CONTENT_TOO_LONG"
    assert runtime.calls == []


def test_rejects_standard_byte_limit_before_runtime() -> None:
    draft = make_docx()
    service, runtime = make_service(max_upload_bytes=len(draft))
    with pytest.raises(AppError) as raised:
        asyncio.run(service.review_document(
            "草稿.docx", draft, review_mode="standard_document",
            standard_filename="标准.docx", standard_content=b"x" * (len(draft) + 1),
        ))
    assert raised.value.code == "UPLOAD_TOO_LARGE"
    assert runtime.calls == []


@pytest.mark.parametrize("kwargs,expected", [
    ({"review_mode": "standard_document"}, "STANDARD_DOCUMENT_REQUIRED"),
    ({"standard_filename": "标准.docx", "standard_content": b"x"}, "INVALID_REVIEW_INPUT"),
    ({"review_mode": "unknown"}, "INVALID_REVIEW_INPUT"),
])
def test_rejects_conflicting_mode_inputs_before_runtime(kwargs, expected) -> None:
    service, runtime = make_service()
    with pytest.raises(AppError) as raised:
        asyncio.run(service.review_document("草稿.docx", make_docx(), **kwargs))
    assert raised.value.code == expected
    assert raised.value.status_code == 422
    assert runtime.calls == []
