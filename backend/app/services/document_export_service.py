"""协调 Markdown 转换、Word 排版及导出结果。"""

import asyncio
import json
import logging
import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import ValidationError

from app.core.config import Settings
from app.core.errors import AppError
from app.infrastructure.pandoc import PandocRunner
from app.schemas.document_export import DocumentExportRequest, ExportedDocument

logger = logging.getLogger(__name__)
EXPORT_QUEUE_TIMEOUT_SECONDS = 5.0


def _prepare_ast(ast: dict[str, object]) -> None:
    if not isinstance(ast, dict) or not (
        isinstance(ast.get("blocks"), list) and isinstance(ast.get("meta"), dict)
        and isinstance(ast.get("pandoc-api-version"), list)
    ):
        raise AppError("DOCUMENT_EXPORT_FAILED", "文档转换结果无效", 500)
    pending = [ast]
    while pending:
        node = pending.pop()
        if isinstance(node, dict):
            content = node.get("c")
            if (
                node.get("t") == "RawInline"
                and isinstance(content, list) and len(content) == 2
                and content[0] == "html" and isinstance(content[1], str)
                and re.fullmatch(r"<br\s*/?>", content[1], flags=re.IGNORECASE)
            ):
                # GFM 表格使用 HTML 换行；转成 Pandoc 原生节点，代码文字不受影响。
                node.clear()
                node["t"] = "LineBreak"
                continue
            if node.get("t") in {"Image", "RawBlock", "RawInline"}:
                raise AppError("EXPORT_CONTENT_UNSUPPORTED", "暂不支持导出图片或原始 HTML", 422)
            pending.extend(node.values())
        elif isinstance(node, list):
            pending.extend(node)


def _build_filename(title: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]', "", title).strip().rstrip(".")
    if name.casefold().endswith(".docx"):
        name = name[:-5]
    name = name[:120].rstrip(" .")
    if not name or re.fullmatch(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])", name.split(".")[0], re.I):
        name = "review-report"
    return f"{name}.docx"


def _font(target, *, size: float, chinese: str = "宋体", latin: str = "Times New Roman"):
    target.font.name = latin
    target.font.size = Pt(size)
    target.font.color.rgb = RGBColor(0, 0, 0)
    target._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), chinese)


def _format_docx(content: bytes) -> bytes:
    document = Document(BytesIO(content))
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in document.styles:
            style = document.styles[name]
            _font(style, size=11)
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
    title = document.styles["Title"]
    _font(title, size=18)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    for level, size in enumerate((16, 14, 12, 11, 11, 11), 1):
        style = document.styles[f"Heading {level}"]
        _font(style, size=size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    for name in ("Source Code", "Verbatim Char"):
        if name in document.styles:
            _font(document.styles[name], size=10, chinese="宋体", latin="Consolas")
    if document.paragraphs:
        document.paragraphs[0].style = title
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Title" or paragraph.style.name.startswith("Heading "):
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
    available = document.sections[0].page_width - document.sections[0].left_margin \
        - document.sections[0].right_margin
    for table in document.tables:
        _format_table(table, available)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _format_table(table, available: int) -> None:
    table.autofit = False
    weights = [max(8, min(72, max((len(cell.text) for cell in column.cells), default=8)))
               for column in table.columns]
    widths = [int(available * weight / sum(weights)) for weight in weights]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn(f"w:{key}"), value)
    for index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        for height in list(row_properties.findall(qn("w:trHeight"))):
            row_properties.remove(height)
        if index == 0 and row_properties.find(qn("w:tblHeader")) is None:
            row_properties.append(OxmlElement("w:tblHeader"))
        for col, cell in enumerate(row.cells):
            cell.width = widths[min(col, len(widths) - 1)]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            for no_wrap in list(cell_properties.findall(qn("w:noWrap"))):
                cell_properties.remove(no_wrap)
            if index == 0:
                shade = cell_properties.find(qn("w:shd"))
                if shade is None:
                    shade = OxmlElement("w:shd")
                    cell_properties.append(shade)
                shade.set(qn("w:fill"), "F2F2F2")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    code = run.style.name in {"Verbatim Char", "Source Code"}
                    _font(run, size=10, latin="Consolas" if code else "Times New Roman")


def _leading_title(ast: dict, title: str) -> bool:
    blocks = ast["blocks"]
    if not blocks or blocks[0].get("t") != "Header":
        return False
    pending = list(reversed(blocks[0].get("c", [None, None, []])[2]))
    parts = []
    while pending:
        inline = pending.pop()
        kind, content = inline.get("t"), inline.get("c")
        if kind == "Str":
            parts.append(content)
        elif kind in {"Space", "SoftBreak", "LineBreak"}:
            parts.append(" ")
        elif kind in {"Strong", "Emph", "Underline", "Strikeout", "SmallCaps",
                      "Superscript", "Subscript"}:
            pending.extend(reversed(content))
        elif kind in {"Link", "Span"}:
            pending.extend(reversed(content[1]))
        elif kind == "Code":
            parts.append(content[1])
        else:
            return False
    return "".join(parts) == title


class DocumentExportService:
    """生成完整文件，供下载 API 和未来业务工具复用。"""

    def __init__(self, settings: Settings, runner: PandocRunner) -> None:
        self._settings = settings
        self._runner = runner
        self._available = False
        self._semaphore = asyncio.Semaphore(settings.document_export_concurrency)
        self._active: set[asyncio.Task] = set()
        self._closed = False
        self._close_task: asyncio.Task | None = None

    async def initialize(self) -> None:
        try:
            await self._runner.initialize()
            self._available = True
        except AppError:
            logger.warning("document_export_unavailable code=DOCUMENT_EXPORT_UNAVAILABLE")

    async def export_docx(
        self, markdown: str, *, title: str = "制度审核报告",
    ) -> ExportedDocument:
        try:
            validated = DocumentExportRequest(markdown=markdown, title=title)
        except ValidationError as exc:
            raise AppError("INVALID_EXPORT_INPUT", "导出内容或标题无效", 422) from exc
        if len(validated.markdown) > self._settings.document_export_max_chars:
            raise AppError("INVALID_EXPORT_INPUT", "导出正文超过字符限制", 422)
        if not self._available or self._closed:
            raise AppError("DOCUMENT_EXPORT_UNAVAILABLE", "文档导出服务不可用", 503)
        task = asyncio.current_task()
        self._active.add(task)
        acquired = False
        try:
            try:
                async with asyncio.timeout(EXPORT_QUEUE_TIMEOUT_SECONDS):
                    await self._semaphore.acquire()
                    acquired = True
            except TimeoutError as exc:
                raise AppError("DOCUMENT_EXPORT_BUSY", "文档导出繁忙，请稍后重试", 503) from exc
            if self._closed:
                raise AppError("DOCUMENT_EXPORT_UNAVAILABLE", "文档导出服务不可用", 503)
            return await self._convert(markdown, title)
        finally:
            if acquired:
                self._semaphore.release()
            self._active.discard(task)

    async def _convert(self, markdown: str, title: str) -> ExportedDocument:
        try:
            deadline = asyncio.get_running_loop().time() \
                + self._settings.document_export_timeout_seconds
            raw_ast = await self._runner.run(
                markdown.encode("utf-8"), source_format="gfm", target_format="json",
                deadline=deadline,
            )
            ast = json.loads(raw_ast)
            _prepare_ast(ast)
            if not _leading_title(ast, title):
                ast["meta"]["title"] = {"t": "MetaString", "c": title}
            else:
                ast["meta"].pop("title", None)
            content = await self._runner.run(
                json.dumps(ast, ensure_ascii=False).encode("utf-8"),
                source_format="json", target_format="docx", deadline=deadline,
            )
            result = await self._wait_work(asyncio.create_task(
                asyncio.to_thread(_format_docx, content),
            ))
            if len(result) > self._settings.document_export_max_output_bytes:
                raise AppError("EXPORT_OUTPUT_TOO_LARGE", "转换文档超过大小限制", 422)
            return ExportedDocument(result, _build_filename(title))
        except AppError:
            raise
        except Exception as exc:
            raise AppError("DOCUMENT_EXPORT_FAILED", "文档转换失败", 500) from exc

    @staticmethod
    async def _wait_work(task: asyncio.Task):
        cancelled = False
        while not task.done():
            try:
                await asyncio.shield(task)
            except asyncio.CancelledError:
                cancelled = True
            except Exception:
                break
        if cancelled:
            # 消费线程异常，但调用方取消始终优先，且槽位保持到线程实际结束。
            if not task.cancelled():
                task.exception()
            raise asyncio.CancelledError
        return task.result()

    async def aclose(self) -> None:
        if self._close_task is None:
            self._closed = True
            self._available = False
            self._close_task = asyncio.create_task(self._close())
        await self._wait_work(self._close_task)

    async def _close(self) -> None:
        tasks = list(self._active)
        for task in tasks:
            task.cancel()
        await asyncio.gather(*tasks, return_exceptions=True)
        await self._runner.aclose()
